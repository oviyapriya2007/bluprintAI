import tempfile
import unittest
from pathlib import Path

import openpyxl
import pandas as pd

from intelligence import build_workspace, export_to_csv, export_to_excel
from intelligence.sample_data import SAMPLE_BOM, SAMPLE_CALLOUTS


class TestExport(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.workspace = build_workspace(bom_data=SAMPLE_BOM, callout_data=SAMPLE_CALLOUTS)

    def test_excel_export(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "out.xlsx"
            result_path = export_to_excel(self.workspace, str(path))
            self.assertTrue(Path(result_path).exists())

            workbook = openpyxl.load_workbook(result_path)
            self.assertEqual(
                workbook.sheetnames, ["BOM", "Procurement Estimate", "AI Validation Issues"]
            )
            bom_sheet = workbook["BOM"]
            self.assertEqual(bom_sheet["A1"].value, "Item Number")
            # header + one row per component
            self.assertEqual(bom_sheet.max_row, len(self.workspace["components"]) + 1)
            self.assertEqual(bom_sheet.freeze_panes, "A2")

    def test_csv_export(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "out.csv"
            result_path = export_to_csv(self.workspace, str(path))
            self.assertTrue(Path(result_path).exists())

            frame = pd.read_csv(result_path)
            self.assertEqual(len(frame), len(self.workspace["components"]))
            self.assertIn("Item Number", frame.columns)
            self.assertIn("Estimated Total (₹)", frame.columns)

    def test_excel_export_with_no_components_does_not_crash(self):
        empty_workspace = build_workspace(bom_data=[], callout_data=[])
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "empty.xlsx"
            export_to_excel(empty_workspace, str(path))
            self.assertTrue(path.exists())


if __name__ == "__main__":
    unittest.main()
